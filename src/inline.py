"""Инлайн-разметка: bold, italic, code, ссылки, sub/sup.

Синтаксис:
    **bold**        -> жирный
    __bold__        -> жирный
    *italic*        -> курсив
    `code`          -> Courier New
    [text](url)     -> гиперссылка
    _{...}          -> нижний индекс
    ^{...}          -> верхний индекс
"""
import re

from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


_TOKENS = re.compile(
    r"(\*\*.+?\*\*"
    r"|__.+?__"
    r"|\*[^*\n]+?\*"
    r"|`[^`\n]+?`"
    r"|\[[^\]\n]+?\]\([^)\n]+?\)"
    r"|\^\{[^}\n]+\}"
    r"|_\{[^}\n]+\}"
    r")"
)


def set_run_font(run, name="Times New Roman", size=14, bold=None, italic=None):
    """Настраивает шрифт, размер, начертание у run'а."""
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)


def add_hyperlink(paragraph, url, text):
    """Добавляет в абзац гиперссылку."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), "Times New Roman")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "28"); rpr.append(sz)
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0000EE"); rpr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    new_run.append(rpr)

    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hl.append(new_run)
    paragraph._p.append(hl)


def _emit(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run_font(run, **kwargs)


def _emit_sup(paragraph, text):
    r = paragraph.add_run(text)
    set_run_font(r)
    r.font.superscript = True


def _emit_sub(paragraph, text):
    r = paragraph.add_run(text)
    set_run_font(r)
    r.font.subscript = True


def add_inline(paragraph, text):
    """Разбирает инлайн-разметку и добавляет runs в абзац."""
    for part in _TOKENS.split(text):
        if not part:
            continue

        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            _emit(paragraph, part[2:-2], bold=True)
        elif part.startswith("__") and part.endswith("__") and len(part) > 4:
            _emit(paragraph, part[2:-2], bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _emit(paragraph, part[1:-1], italic=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            _emit(paragraph, part[1:-1], name="Courier New", size=12)
        elif part.startswith("^{") and part.endswith("}"):
            _emit_sup(paragraph, part[2:-1])
        elif part.startswith("_{") and part.endswith("}"):
            _emit_sub(paragraph, part[2:-1])
        else:
            m = re.match(r"^\[([^\]]+)\]\(([^)]+)\)$", part)
            if m:
                add_hyperlink(paragraph, m.group(2), m.group(1))
            else:
                _emit(paragraph, part)
