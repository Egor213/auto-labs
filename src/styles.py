"""Настройка страницы, стилей и нумерации."""
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_style_font(style, name="Times New Roman", size=14, bold=None):
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)


def _get_or_create_paragraph_style(doc, name, base="Normal"):
    """Возвращает стиль, создавая его при отсутствии."""
    styles = doc.styles
    try:
        return styles[name]
    except KeyError:
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        try:
            st.base_style = styles[base]
        except KeyError:
            pass
        return st


def setup_page(doc):
    """Поля: лево 3 см, право 1.5 см, верх/низ 2 см. A4."""
    s = doc.sections[0]
    s.page_width = Mm(210)
    s.page_height = Mm(297)
    s.top_margin = Cm(2)
    s.bottom_margin = Cm(2)
    s.left_margin = Cm(3)
    s.right_margin = Cm(1.5)


def setup_styles(doc):
    styles = doc.styles

    # --- Normal -----------------------------------------------------
    normal = styles["Normal"]
    _set_style_font(normal, size=14)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    # --- Heading 1/2/3 (создаём при отсутствии) ---------------------
    h1 = _get_or_create_paragraph_style(doc, "Heading 1")
    _set_style_font(h1, size=14, bold=True)
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.font.all_caps = True
    pf = h1.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.keep_with_next = True

    h2 = _get_or_create_paragraph_style(doc, "Heading 2")
    _set_style_font(h2, size=14, bold=True)
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.font.all_caps = False
    pf = h2.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25)
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.keep_with_next = True

    h3 = _get_or_create_paragraph_style(doc, "Heading 3")
    _set_style_font(h3, size=14, bold=True)
    h3.font.color.rgb = RGBColor(0, 0, 0)
    pf = h3.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25)
    pf.keep_with_next = True

    # --- Листинг ----------------------------------------------------
    listing = _get_or_create_paragraph_style(doc, "Listing")
    _set_style_font(listing, name="Courier New", size=11)
    pf = listing.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.0
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    # --- Списки -----------------------------------------------------
    for name in ("List Bullet", "List Number"):
        try:
            st = styles[name]
        except KeyError:
            continue
        _set_style_font(st, size=14)
        pf = st.paragraph_format
        pf.line_spacing = 1.5
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def setup_page_numbers(doc):
    """Номер страницы внизу справа, без номера на титульном листе."""
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Cm(0)

    run = p.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)