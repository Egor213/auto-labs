#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Bootstrap проекта md2docx.
Создаёт структуру:

    project/
      ├── run.py
      ├── requirements.txt
      ├── README.md
      ├── examples/example.md
      └── src/
          ├── __init__.py
          ├── config.py
          ├── template.yaml
          ├── frontmatter.py
          ├── inline.py
          ├── styles.py
          ├── elements.py
          ├── title_page.py
          ├── parser.py
          └── md2docx.py

Запуск:
    python create_project.py            # в текущую папку
    python create_project.py --force    # перезаписать существующие
    python create_project.py -o ./proj  # куда создавать
"""

from __future__ import annotations

import argparse
from pathlib import Path


FILES: dict[str, str] = {}

# -------------------------------------------------------------------------
# src/__init__.py
# -------------------------------------------------------------------------
FILES["src/__init__.py"] = r'''"""md2docx - конвертер Markdown в DOCX для отчетов ЛЭТИ."""
from .md2docx import convert, main

__all__ = ["convert", "main"]
'''


# -------------------------------------------------------------------------
# src/template.yaml
# -------------------------------------------------------------------------
FILES["src/template.yaml"] = r'''# Значения по умолчанию для титульного листа.
# Могут быть переопределены в YAML frontmatter конкретного .md файла.
university: |
  САНКТ-ПЕТЕРБУРГСКИЙ ГОСУДАРСТВЕННЫЙ
  ЭЛЕКТРОТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ
  «ЛЭТИ» ИМ. В.И. УЛЬЯНОВА (ЛЕНИНА)
ministry: МИНОБРНАУКИ РОССИИ
department: МО ЭВМ
city: Санкт-Петербург
year: "2024"
'''


# -------------------------------------------------------------------------
# src/config.py
# -------------------------------------------------------------------------
FILES["src/config.py"] = r'''"""Загрузка значений по умолчанию из template.yaml."""
from pathlib import Path

try:
    import yaml
except ImportError:
    yaml = None  # type: ignore

_HERE = Path(__file__).parent
_TEMPLATE = _HERE / "template.yaml"

FALLBACK = {
    "ministry": "МИНОБРНАУКИ РОССИИ",
    "university": (
        "САНКТ-ПЕТЕРБУРГСКИЙ ГОСУДАРСТВЕННЫЙ\n"
        "ЭЛЕКТРОТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ\n"
        "«ЛЭТИ» ИМ. В.И. УЛЬЯНОВА (ЛЕНИНА)"
    ),
    "department": "МО ЭВМ",
    "city": "Санкт-Петербург",
    "year": "2024",
}


def load_defaults() -> dict:
    """Возвращает словарь значений по умолчанию из template.yaml."""
    if yaml is None or not _TEMPLATE.exists():
        return dict(FALLBACK)
    try:
        with _TEMPLATE.open(encoding="utf-8") as f:
            data = yaml.safe_load(f) or {}
    except Exception:
        return dict(FALLBACK)
    merged = dict(FALLBACK)
    for k, v in data.items():
        if v is not None:
            merged[k] = v
    return merged
'''


# -------------------------------------------------------------------------
# src/frontmatter.py
# -------------------------------------------------------------------------
FILES["src/frontmatter.py"] = r'''"""Разбор YAML frontmatter и слияние с template.yaml."""
import re

try:
    import yaml
except ImportError:
    yaml = None  # type: ignore

from .config import load_defaults

_FM_RE = re.compile(r"^---\s*\n(.*?)\n---\s*\n?(.*)$", re.DOTALL)


def parse(text: str):
    """Возвращает (meta, body)."""
    text = text.lstrip("\ufeff")
    defaults = load_defaults()

    if not text.startswith("---"):
        return defaults, text

    m = _FM_RE.match(text)
    if not m:
        return defaults, text

    meta = {}
    if yaml is not None:
        try:
            meta = yaml.safe_load(m.group(1)) or {}
        except Exception:
            meta = {}

    merged = dict(defaults)
    for k, v in meta.items():
        if v is not None:
            merged[k] = v
    return merged, m.group(2)
'''


# -------------------------------------------------------------------------
# src/inline.py
# -------------------------------------------------------------------------
FILES["src/inline.py"] = r'''"""Инлайн-разметка: bold, italic, code, ссылки, sub/sup.

Синтаксис:
    **bold**        -> жирный
    __bold__        -> жирный
    *italic*        -> курсив
    `code`          -> Courier New
    [text](url)     -> гиперссылка
    _{...}          -> нижний индекс
    ^{...}          -> верхний индекс
"""
import re

from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


_TOKENS = re.compile(
    r"(\*\*.+?\*\*"
    r"|__.+?__"
    r"|\*[^*\n]+?\*"
    r"|`[^`\n]+?`"
    r"|\[[^\]\n]+?\]\([^)\n]+?\)"
    r"|\^\{[^}\n]+\}"
    r"|_\{[^}\n]+\}"
    r")"
)


def set_run_font(run, name="Times New Roman", size=14, bold=None, italic=None):
    """Настраивает шрифт, размер, начертание у run'а."""
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)


def add_hyperlink(paragraph, url, text):
    """Добавляет в абзац гиперссылку."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(a), "Times New Roman")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "28"); rpr.append(sz)
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0000EE"); rpr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    new_run.append(rpr)

    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hl.append(new_run)
    paragraph._p.append(hl)


def _emit(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run_font(run, **kwargs)


def _emit_sup(paragraph, text):
    r = paragraph.add_run(text)
    set_run_font(r)
    r.font.superscript = True


def _emit_sub(paragraph, text):
    r = paragraph.add_run(text)
    set_run_font(r)
    r.font.subscript = True


def add_inline(paragraph, text):
    """Разбирает инлайн-разметку и добавляет runs в абзац."""
    for part in _TOKENS.split(text):
        if not part:
            continue

        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            _emit(paragraph, part[2:-2], bold=True)
        elif part.startswith("__") and part.endswith("__") and len(part) > 4:
            _emit(paragraph, part[2:-2], bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _emit(paragraph, part[1:-1], italic=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            _emit(paragraph, part[1:-1], name="Courier New", size=12)
        elif part.startswith("^{") and part.endswith("}"):
            _emit_sup(paragraph, part[2:-1])
        elif part.startswith("_{") and part.endswith("}"):
            _emit_sub(paragraph, part[2:-1])
        else:
            m = re.match(r"^\[([^\]]+)\]\(([^)]+)\)$", part)
            if m:
                add_hyperlink(paragraph, m.group(2), m.group(1))
            else:
                _emit(paragraph, part)
'''


# -------------------------------------------------------------------------
# src/styles.py
# -------------------------------------------------------------------------
FILES["src/styles.py"] = r'''"""Настройка страницы, стилей и нумерации."""
from docx.shared import Pt, Cm, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_style_font(style, name="Times New Roman", size=14, bold=None):
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)


def setup_page(doc):
    """Поля: лево 3 см, право 1.5 см, верх/низ 2 см. A4."""
    s = doc.sections[0]
    s.page_width = Mm(210)
    s.page_height = Mm(297)
    s.top_margin = Cm(2)
    s.bottom_margin = Cm(2)
    s.left_margin = Cm(3)
    s.right_margin = Cm(1.5)


def setup_styles(doc):
    styles = doc.styles

    # --- Основной текст ---
    normal = styles["Normal"]
    _set_style_font(normal, size=14)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    # --- Заголовок 1 ---
    h1 = styles["Heading 1"]
    _set_style_font(h1, size=14, bold=True)
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.font.all_caps = True
    pf = h1.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.keep_with_next = True

    # --- Заголовок 2 ---
    h2 = styles["Heading 2"]
    _set_style_font(h2, size=14, bold=True)
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.font.all_caps = False
    pf = h2.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(1.25)
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.keep_with_next = True

    # --- Заголовок 3 ---
    try:
        h3 = styles["Heading 3"]
        _set_style_font(h3, size=14, bold=True)
        h3.font.color.rgb = RGBColor(0, 0, 0)
        pf = h3.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing = 1.5
        pf.first_line_indent = Cm(1.25)
        pf.keep_with_next = True
    except KeyError:
        pass

    # --- Листинг ---
    try:
        listing = styles["Listing"]
    except KeyError:
        listing = styles.add_style("Listing", WD_STYLE_TYPE.PARAGRAPH)
        listing.base_style = styles["Normal"]
    _set_style_font(listing, name="Courier New", size=11)
    pf = listing.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.0
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    # --- Списки ---
    for name in ("List Bullet", "List Number"):
        try:
            st = styles[name]
            _set_style_font(st, size=14)
            pf = st.paragraph_format
            pf.line_spacing = 1.5
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        except KeyError:
            pass


def setup_page_numbers(doc):
    """Номер страницы внизу справа, без номера на титульном листе."""
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Cm(0)

    run = p.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
'''


# -------------------------------------------------------------------------
# src/elements.py
# -------------------------------------------------------------------------
FILES["src/elements.py"] = r'''"""Функции добавления элементов в документ."""
import re
from pathlib import Path

from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .inline import add_inline, set_run_font

try:
    from PIL import Image
    _HAS_PIL = True
except ImportError:
    _HAS_PIL = False


# ---------------------------------------------------------------------
#  Заголовки, абзацы, списки
# ---------------------------------------------------------------------
def add_heading(doc, text, level):
    level = max(1, min(level, 3))
    p = doc.add_paragraph(style="Heading {0}".format(level))
    add_inline(p, text)
    return p


def add_paragraph_multiline(doc, lines):
    """Один абзац, строки внутри разделены soft-break (<w:br/>)."""
    p = doc.add_paragraph()
    for idx, line in enumerate(lines):
        if idx > 0:
            p.add_run().add_break()
        add_inline(p, line)
    return p


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    add_inline(p, text)
    return p


def add_list_item(doc, text, ordered=False):
    style = "List Number" if ordered else "List Bullet"
    try:
        p = doc.add_paragraph(style=style)
    except KeyError:
        p = doc.add_paragraph()
    add_inline(p, text)
    return p


# ---------------------------------------------------------------------
#  Подписи
# ---------------------------------------------------------------------
def add_caption_above(doc, caption, align=None):
    p = doc.add_paragraph()
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.line_spacing = 1.0
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)
    run = p.add_run(caption)
    set_run_font(run, size=14)
    return p


def add_caption_below(doc, caption, align=None):
    p = doc.add_paragraph()
    p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.line_spacing = 1.0
    pf.space_before = Pt(3)
    pf.space_after = Pt(6)
    run = p.add_run(caption)
    set_run_font(run, size=14)
    return p


# ---------------------------------------------------------------------
#  Листинг
# ---------------------------------------------------------------------
def add_code_block(doc, code, caption=None):
    if caption:
        add_caption_above(doc, caption)
    for line in code.split("\n"):
        p = doc.add_paragraph(style="Listing")
        run = p.add_run(line if line else " ")
        set_run_font(run, name="Courier New", size=11)
    tail = doc.add_paragraph()
    tail.paragraph_format.first_line_indent = Cm(0)
    tail.paragraph_format.line_spacing = 1.0
    tail.paragraph_format.space_after = Pt(0)


# ---------------------------------------------------------------------
#  Рисунок
# ---------------------------------------------------------------------
def add_image(doc, path, caption=""):
    path = Path(path)
    if not path.exists():
        print("[WARN] Изображение не найдено: {0}".format(path))
        return

    width_cm = 15.0
    if _HAS_PIL:
        try:
            with Image.open(path) as im:
                w, h = im.size
                dpi = im.info.get("dpi", (96, 96))[0] or 96
                native_cm = w / dpi * 2.54
                width_cm = min(native_cm, 16.0)
        except Exception:
            pass

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))

    if caption:
        add_caption_below(doc, caption)


# ---------------------------------------------------------------------
#  Таблица
# ---------------------------------------------------------------------
_SEP = re.compile(r"^:?-+:?$")


def _parse_table_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def add_table(doc, table_lines, caption=None):
    if caption:
        add_caption_above(doc, caption)

    rows = [_parse_table_row(l) for l in table_lines]
    if not rows:
        return

    header = rows[0]
    data = []
    for r in rows[1:]:
        if r and all(_SEP.match(c) for c in r if c):
            continue
        data.append(r)

    n_cols = len(header)
    fs = 12 if len(data) > 5 else 14

    table = doc.add_table(rows=1 + len(data), cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for j, txt in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(txt)
        set_run_font(run, size=fs, bold=True)

    for i, row in enumerate(data):
        for j in range(n_cols):
            txt = row[j] if j < len(row) else ""
            cell = table.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(txt)
            set_run_font(run, size=fs)

    tail = doc.add_paragraph()
    tail.paragraph_format.first_line_indent = Cm(0)
    tail.paragraph_format.line_spacing = 1.0
    tail.paragraph_format.space_after = Pt(0)


# ---------------------------------------------------------------------
#  Формула
# ---------------------------------------------------------------------
def add_formula(doc, formula):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    add_inline(p, formula)
'''


# -------------------------------------------------------------------------
# src/title_page.py
# -------------------------------------------------------------------------
FILES["src/title_page.py"] = r'''"""Титульный лист по шаблону ЛЭТИ."""
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT

from .inline import set_run_font


def _para(doc, text="", size=14, bold=False,
          align="center", indent_cm=0.0, space_after=0):
    p = doc.add_paragraph()
    amap = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }
    p.alignment = amap[align]
    pf = p.paragraph_format
    pf.first_line_indent = Cm(indent_cm)
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold)
    return p


def add_title_page(doc, meta):
    """Собирает титульный лист по шаблону кафедры МО ЭВМ."""

    # -------- Шапка ----------------------------------------------------
    _para(doc, meta.get("ministry", "МИНОБРНАУКИ РОССИИ"), size=12)

    uni = meta.get("university", "") or ""
    for line in uni.split("\n"):
        line = line.strip()
        if line:
            _para(doc, line, size=12)

    dept = meta.get("department", "")
    if dept:
        _para(doc, "Кафедра {0}".format(dept), size=12)

    # -------- Отступ --------------------------------------------------
    for _ in range(4):
        _para(doc, "")

    # -------- Заголовок отчёта ---------------------------------------
    _para(doc, "ОТЧЕТ", size=14, bold=True)
    _para(doc, "по лабораторной работе № {0}".format(
        meta.get("lab_number", "")), size=14)
    _para(doc, "по дисциплине «{0}»".format(
        meta.get("discipline", "")), size=14)
    _para(doc, "Тема: {0}".format(meta.get("topic", "")), size=14)

    # -------- Отступ --------------------------------------------------
    for _ in range(3):
        _para(doc, "")

    # -------- Студент / Преподаватель --------------------------------
    # Две колонки через tab stop примерно на середине текста.
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8.0),
                                              WD_TAB_ALIGNMENT.LEFT)
    r = p.add_run("Студент(ка) гр. {0}\tПреподаватель".format(
        meta.get("group", "")))
    set_run_font(r, size=14)

    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8.0),
                                              WD_TAB_ALIGNMENT.LEFT)
    r = p.add_run("{0}\t{1}".format(
        meta.get("student", ""), meta.get("teacher", "")))
    set_run_font(r, size=14)

    # -------- Отступ --------------------------------------------------
    for _ in range(4):
        _para(doc, "")

    # -------- Город / Год --------------------------------------------
    _para(doc, meta.get("city", "Санкт-Петербург"), size=14)
    _para(doc, str(meta.get("year", "")), size=14)

    doc.add_page_break()
'''


# -------------------------------------------------------------------------
# src/parser.py
# -------------------------------------------------------------------------
FILES["src/parser.py"] = r'''"""Разбор тела Markdown-документа.

Особенности:
    * блок подряд идущих непустых "обычных" строк склеивается в ОДИН
      абзац Word, а переходы между строками становятся soft-break
      (<w:br/>), то есть Shift+Enter в Word;
    * пустая строка = конец абзаца;
    * спец. элементы (заголовки, листинги, таблицы, рисунки, списки,
      формулы) разбираются построчно, как раньше.
"""
import re
from pathlib import Path

from .elements import (
    add_heading, add_paragraph_multiline, add_list_item, add_code_block,
    add_image, add_table, add_formula,
)

_RE_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_RE_IMG = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
_RE_BULLET = re.compile(r"^[-*+]\s+(.*)$")
_RE_ORDERED = re.compile(r"^\d+\.\s+(.*)$")
_RE_TABLE_CAP = re.compile(r"^Таблица\s+\d+\s*-", re.IGNORECASE)
_RE_LISTING_CAP = re.compile(r"^Листинг\s+\d+\s*-", re.IGNORECASE)
_RE_TABLE_SEP = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?\s*$")


def _is_special_start(s: str) -> bool:
    """Начало ли это спец-элемента (не часть обычного абзаца)."""
    if not s:
        return True
    if s.startswith("```"):
        return True
    if s.startswith("$$"):
        return True
    if _RE_IMG.match(s):
        return True
    if _RE_TABLE_CAP.match(s):
        return True
    if _RE_HEADING.match(s):
        return True
    if _RE_BULLET.match(s):
        return True
    if _RE_ORDERED.match(s):
        return True
    return False


def _is_table_start(lines, i):
    """Строка i начинает таблицу (pipe + строка-разделитель ниже)."""
    s = lines[i].strip()
    if "|" not in s:
        return False
    return i + 1 < len(lines) and bool(_RE_TABLE_SEP.match(lines[i + 1].strip()))


def process_body(doc, body, base_dir):
    lines = body.split("\n")
    n = len(lines)
    i = 0

    while i < n:
        stripped = lines[i].strip()

        # пустая строка
        if not stripped:
            i += 1
            continue

        # ``` - блок кода / листинг
        if stripped.startswith("```"):
            info = stripped[3:].strip()
            caption = None
            if ":" in info:
                _, caption = info.split(":", 1)
                caption = caption.strip()
            if caption is None and i > 0:
                prev = lines[i - 1].strip()
                if _RE_LISTING_CAP.match(prev):
                    caption = prev

            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # пропускаем закрывающий ```
            add_code_block(doc, "\n".join(code_lines), caption)
            continue

        # $$ - формула (display)
        if stripped.startswith("$$"):
            if stripped.endswith("$$") and len(stripped) > 4:
                add_formula(doc, stripped[2:-2].strip())
                i += 1
                continue
            buf = [stripped[2:]]
            i += 1
            while i < n and "$$" not in lines[i]:
                buf.append(lines[i])
                i += 1
            if i < n:
                buf.append(lines[i].split("$$")[0])
                i += 1
            add_formula(doc, "\n".join(buf).strip())
            continue

        # ![Рисунок N -- Имя](path)
        m = _RE_IMG.match(stripped)
        if m:
            caption = m.group(1)
            src = m.group(2)
            add_image(doc, (Path(base_dir) / src).resolve(), caption)
            i += 1
            continue

        # "Таблица N -- Имя" + pipe-таблица
        if _RE_TABLE_CAP.match(stripped):
            j = i + 1
            while j < n and not lines[j].strip():
                j += 1
            if j < n and "|" in lines[j]:
                caption = stripped
                tlines = []
                while j < n and "|" in lines[j]:
                    tlines.append(lines[j])
                    j += 1
                add_table(doc, tlines, caption)
                i = j
                continue

        # pipe-таблица без подписи
        if _is_table_start(lines, i):
            tlines = []
            j = i
            while j < n and "|" in lines[j]:
                tlines.append(lines[j])
                j += 1
            add_table(doc, tlines, None)
            i = j
            continue

        # # Заголовок
        m = _RE_HEADING.match(stripped)
        if m:
            add_heading(doc, m.group(2).strip(), len(m.group(1)))
            i += 1
            continue

        # - пункт
        m = _RE_BULLET.match(stripped)
        if m:
            add_list_item(doc, m.group(1), ordered=False)
            i += 1
            continue

        # 1. пункт
        m = _RE_ORDERED.match(stripped)
        if m:
            add_list_item(doc, m.group(1), ordered=True)
            i += 1
            continue

        # обычный абзац: собираем подряд идущие непустые "обычные" строки
        buf = [stripped]
        i += 1
        while i < n:
            nxt = lines[i].strip()
            if not nxt:
                break
            if _is_special_start(nxt):
                break
            if _is_table_start(lines, i):
                break
            buf.append(nxt)
            i += 1
        add_paragraph_multiline(doc, buf)
'''


# -------------------------------------------------------------------------
# src/md2docx.py
# -------------------------------------------------------------------------
FILES["src/md2docx.py"] = r'''"""Основной конвертер Markdown -> DOCX."""
import argparse
import sys
from pathlib import Path

from docx import Document

from .frontmatter import parse as parse_frontmatter
from .styles import setup_page, setup_styles, setup_page_numbers
from .title_page import add_title_page
from .parser import process_body


def convert(md_path, docx_path=None):
    """Конвертирует .md в .docx. Возвращает путь к созданному файлу."""
    md_path = Path(md_path)
    if not md_path.exists():
        raise FileNotFoundError(md_path)

    text = md_path.read_text(encoding="utf-8")
    meta, body = parse_frontmatter(text)

    doc = Document()
    setup_page(doc)
    setup_styles(doc)
    setup_page_numbers(doc)
    add_title_page(doc, meta)
    process_body(doc, body, md_path.parent)

    if docx_path is None:
        docx_path = md_path.with_suffix(".docx")
    docx_path = Path(docx_path)
    doc.save(str(docx_path))
    print("[OK] {0} -> {1}".format(md_path, docx_path))
    return docx_path


def main(argv=None):
    ap = argparse.ArgumentParser(description="MD -> DOCX для отчетов ЛЭТИ")
    ap.add_argument("input", nargs="?", help="Входной .md файл")
    ap.add_argument("output", nargs="?", help="Выходной .docx (опционально)")
    args = ap.parse_args(argv)

    if not args.input:
        ap.print_help()
        sys.exit(1)
    convert(args.input, args.output)


if __name__ == "__main__":
    main()
'''


# -------------------------------------------------------------------------
# run.py
# -------------------------------------------------------------------------
FILES["run.py"] = r'''#!/usr/bin/env python3
"""Точка входа: python run.py report.md [out.docx]"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from src.md2docx import main

if __name__ == "__main__":
    main()
'''


# -------------------------------------------------------------------------
# requirements.txt
# -------------------------------------------------------------------------
FILES["requirements.txt"] = (
    "python-docx>=0.8.11\n"
    "PyYAML>=6.0\n"
    "pillow>=9.0\n"
)


# -------------------------------------------------------------------------
# README.md
# -------------------------------------------------------------------------
FILES["README.md"] = r'''# md2docx

Конвертер Markdown -> DOCX для отчётов по лабораторным работам СПбГЭТУ
«ЛЭТИ» (кафедра МО ЭВМ). Формирует титульный лист по шаблону кафедры,
применяет нужные стили (шрифты, отступы, поля, нумерацию страниц).

## Установка

```bash
pip install -r requirements.txt
```

## Использование

```bash
python run.py examples/example.md
python run.py examples/example.md out.docx
```

## Формат Markdown

В начале файла -- YAML frontmatter:

```yaml
---
university: |
  САНКТ-ПЕТЕРБУРГСКИЙ ГОСУДАРСТВЕННЫЙ
  ЭЛЕКТРОТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ
  «ЛЭТИ» ИМ. В.И. УЛЬЯНОВА (ЛЕНИНА)
department: МО ЭВМ
discipline: Базы данных
lab_number: 4
topic: Нагрузочное тестирование БД
student: Муравин Е. Е.
group: 2382
teacher: Заславский М. М.
city: Санкт-Петербург
year: 2024
---
```

Значения по умолчанию -- в `src/template.yaml`.

### Разметка тела

| Элемент                            | Синтаксис |
|------------------------------------|-----------|
| Заголовок 1 (приложения)           | `# ПРИЛОЖЕНИЕ А` |
| Заголовок 2 (разделы отчёта)       | `## Цель работы` |
| Заголовок 3                        | `### Подраздел` |
| Абзац                              | текст; перенос строки внутри блока = soft break |
| Список                             | `- пункт` / `1. пункт` |
| Жирный, курсив, код                | `**b**`, `*i*`, `` `c` `` |
| Ссылка                             | `[текст](https://url)` |
| Нижний / верхний индекс            | `_{2}` / `^{2}` |
| Формула                            | `$$ ... $$` |
| Рисунок                            | `![Рисунок 1 -- Имя](img.png)` |
| Таблица                            | строка `Таблица 1 -- Имя` + pipe-таблица |
| Листинг                            | ` ```js: Листинг 1 -- Имя ``` ` |

### Правило абзаца

Подряд идущие непустые строки склеиваются в ОДИН абзац Word. Переносы
между ними вставляются как soft-break (`<w:br/>`, Shift+Enter в Word).
Пустая строка разделяет абзацы.

```
Строка 1 блока.
Строка 2 блока.

Новый абзац.
```

даёт один абзац из двух строк и второй абзац из одной.

## Структура

```
src/
  config.py        -- значения по умолчанию
  template.yaml    -- редактируемые дефолты
  frontmatter.py   -- разбор YAML frontmatter
  inline.py        -- инлайн-разметка
  styles.py        -- стили docx
  elements.py      -- добавление элементов
  title_page.py    -- титульный лист
  parser.py        -- парсер тела MD
  md2docx.py       -- точка входа конвертера
run.py             -- CLI-обёртка
examples/          -- примеры
```
'''


# -------------------------------------------------------------------------
# examples/example.md
# -------------------------------------------------------------------------
FILES["examples/example.md"] = r'''---
university: |
  САНКТ-ПЕТЕРБУРГСКИЙ ГОСУДАРСТВЕННЫЙ
  ЭЛЕКТРОТЕХНИЧЕСКИЙ УНИВЕРСИТЕТ
  «ЛЭТИ» ИМ. В.И. УЛЬЯНОВА (ЛЕНИНА)
department: МО ЭВМ
discipline: Базы данных
lab_number: 4
topic: Нагрузочное тестирование БД
student: Муравин Е. Е.
group: 2382
teacher: Заславский М. М.
city: Санкт-Петербург
year: 2024
---

## Цель работы

Ознакомиться со структурой данных **«индексы»** в языке SQL.
Протестировать и замерить время выполнения запросов из предыдущей
лабораторной работы на различных объёмах данных.

## Основные теоретические положения

SQL -- *Structured Query Language* -- стандартный язык запросов для
работы с реляционными базами данных. PostgreSQL -- объектно-реляционная
СУБД с открытым исходным кодом.

Основные возможности SQL:

- создание баз данных и таблиц;
- чтение, вставка, обновление и удаление данных;
- управление индексами и ограничениями;
- агрегация и группировка.

Документация доступна по ссылке: [PostgreSQL 16](https://www.postgresql.org/docs/16/).

## Задание

Вариант 7.

1. Написать скрипт, заполняющий БД большим количеством тестовых данных.
2. Измерить время выполнения запросов для разного числа записей.
3. Добавить в БД индексы (не менее 5 штук).
4. Оценить влияние сортировки на скорость выполнения запросов.

## Выполнение работы

Для заполнения таблиц был создан файл `generate_faker.js`, в котором
данные генерируются с использованием библиотеки **faker.js**.

Структура базы данных приведена на рисунке 1.

![Рисунок 1 -- Структура базы данных](structure.png)

Результаты первого тестирования представлены в таблице 1.

Таблица 1 -- Время выполнения запросов без сортировки

| Количество записей | task1 (сек) | task2 (сек) | task3 (сек) |
|-------------------:|------------:|------------:|------------:|
| 100                | 0.006       | 0.008       | 0.004       |
| 1 000              | 0.005       | 0.007       | 0.004       |
| 10 000             | 0.007       | 0.010       | 0.008       |
| 1 000 000          | 0.152       | 0.246       | 0.164       |

Как видно из результатов, наибольшее время выполнения показали
запросы с группировкой данных.

Пример функции генерации данных приведён в листинге 1.

```javascript: Листинг 1 -- Функция create_data
const { Faker } = require('faker');

async function create_data(count) {
    const diets = Array.from({ length: count },
        () => Faker.getRandomDiet());
    await Diet.bulkCreate(diets);
}
```

Среднее время выполнения одного запроса оценивается формулой:

$$ \bar{t} = (1/n) * sum(t_{i}) $$

Здесь t_{i} -- время выполнения i-го запроса, n -- количество замеров.

## Выводы

В ходе работы была изучена библиотека faker.js для генерации
правдоподобных данных, исследовано влияние индексов на скорость
выполнения запросов. Индексы значительно ускорили выборку, но
замедлили вставку.

# ПРИЛОЖЕНИЕ А

Полный исходный код доступен по ссылке:
[DB Fiddle](https://www.db-fiddle.com/).
'''


# -------------------------------------------------------------------------
#  MAIN
# -------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(
        description="Bootstrap проекта md2docx: создаёт структуру файлов."
    )
    ap.add_argument("-o", "--output", default=".",
                    help="Каталог для создания проекта (по умолчанию: текущий)")
    ap.add_argument("--force", action="store_true",
                    help="Перезаписать существующие файлы")
    args = ap.parse_args()

    root = Path(args.output).resolve()
    root.mkdir(parents=True, exist_ok=True)

    created = skipped = 0
    for rel, content in FILES.items():
        path = root / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        if path.exists() and not args.force:
            print("[SKIP]   {0}".format(rel))
            skipped += 1
            continue
        path.write_text(content.rstrip("\n") + "\n", encoding="utf-8")
        print("[CREATE] {0}".format(rel))
        created += 1

    print()
    print("Готово. Создано: {0}, пропущено: {1}".format(created, skipped))
    print()
    print("Далее:")
    print("  cd {0}".format(root))
    print("  pip install -r requirements.txt")
    print("  python run.py examples/example.md")


if __name__ == "__main__":
    main()
